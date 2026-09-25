import tempfile
import unittest
import sys
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from office_copy import make_document

class OfficeCopyTests(unittest.TestCase):
    def test_paper_styles_and_editable_math(self):
        fragment = '<h2>标题</h2><ul><li>中文 QWIP <b>加粗</b><math xmlns="http://www.w3.org/1998/Math/MathML"><mfrac><mi>x</mi><mi>y</mi></mfrac></math></li></ul><p>下一段</p>'
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'copy.docx'
            make_document(fragment, path)
            d = Document(path)
            normal = d.styles['Normal']
            self.assertEqual(normal.font.name, 'Times New Roman')
            self.assertEqual(normal.font.size.pt, 10.5)
            self.assertEqual(normal.element.rPr.rFonts.get(qn('w:eastAsia')), 'SimSun')
            self.assertEqual(normal.paragraph_format.line_spacing, 1.4)
            self.assertEqual(normal.paragraph_format.first_line_indent.pt, 21)
            self.assertEqual(len(d.element.xpath('.//m:f')), 1)
            self.assertEqual(len(d.element.xpath('.//w:numPr')), 0)
            self.assertEqual(len(d.element.xpath('.//w:drawing')), 0)
            self.assertEqual(len(d.settings.element.findall(qn('m:mathPr'))), 1)
            self.assertIn('下一段', d.paragraphs[-1].text)

    def test_table_and_display_equation(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'copy.docx'
            make_document('<table><tr><th>参数</th><th>值</th></tr><tr><td>x</td><td>1</td></tr></table><div><math xmlns="http://www.w3.org/1998/Math/MathML" display="block"><msqrt><mi>x</mi></msqrt></math></div>', path)
            d = Document(path)
            self.assertEqual(len(d.tables), 1)
            self.assertEqual(d.tables[0].cell(1, 1).text, '1')
            self.assertEqual(d.paragraphs[-1].alignment, 1)
            self.assertEqual(len(d.element.xpath('.//m:rad')), 1)
