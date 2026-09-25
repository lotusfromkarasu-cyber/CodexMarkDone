"""Native Word/WPS clipboard, with paper styles and editable OMML equations."""
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
import tempfile
import threading

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from lxml import etree, html
import mathml2omml

MATH_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
_MATHML_NS = 'http://www.w3.org/1998/Math/MathML'
_MATHML_ARITY = {
    'mfrac': 2,
    'mroot': 2,
    'msub': (2, 3),
    'msup': (2, 3),
    'msubsup': (3, 4),
    'munder': (2, 3),
    'mover': (2, 3),
    'munderover': (3, 4),
}


class _MathMLOperatorView:
    def __init__(self, element):
        self.attrs = element.attrib
        self._text = ''.join(element.itertext())

    def text(self):
        return self._text


def _group_excess_mathml_children(source):
    """Group extra operands before OMML conversion while preserving script arity."""
    parser = etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False)
    root = etree.fromstring(source.encode('utf-8'), parser)
    changed = False
    for node in root.iter():
        if not isinstance(node.tag, str):
            continue
        tag = etree.QName(node).localname.lower()
        arities = _MATHML_ARITY.get(tag)
        if arities is None:
            continue
        children = [
            child for child in node
            if isinstance(child.tag, str)
            and etree.QName(child).localname.lower() not in ('annotation', 'annotation-xml')
        ]
        max_arity = arities[0] if isinstance(arities, tuple) else arities
        if isinstance(arities, tuple):
            first = children[0] if children else None
            if first is not None and etree.QName(first).localname.lower() == 'mo':
                op = mathml2omml.op_attrs(_MathMLOperatorView(first))
                if op['largeop'] and op['form'] == 'prefix':
                    max_arity = arities[1]
        if len(children) <= max_arity:
            continue

        extras = children[max_arity - 1:]
        current_children = list(node)
        insert_at = sum(child not in extras for child in current_children[:current_children.index(extras[0])])
        namespace = etree.QName(node).namespace or _MATHML_NS
        grouped = etree.Element(f'{{{namespace}}}mrow')
        for child in extras:
            node.remove(child)
        node.insert(insert_at, grouped)
        for child in extras:
            grouped.append(child)
        changed = True

    if not changed:
        return source
    return etree.tostring(root, encoding='unicode')


def _hide_square_root_degree_placeholder(omml):
    """Mark radicals without an index as square roots in Office Math."""
    wrapper = etree.fromstring(
        ('<root xmlns:m="' + MATH_NS + '">' + omml + '</root>').encode('utf-8')
    )
    namespaces = {'m': MATH_NS}
    for radical in wrapper.xpath('.//m:rad[not(m:deg)]', namespaces=namespaces):
        properties = radical.find(qn('m:radPr'))
        if properties is None:
            properties = OxmlElement('m:radPr')
            radical.insert(0, properties)
        degree_hidden = properties.find(qn('m:degHide'))
        if degree_hidden is None:
            degree_hidden = OxmlElement('m:degHide')
            properties.append(degree_hidden)
        degree_hidden.set(qn('m:val'), '1')
        if radical.find(qn('m:deg')) is None:
            degree = OxmlElement('m:deg')
            radical.insert(radical.index(properties) + 1, degree)
    return ''.join(etree.tostring(child, encoding='unicode') for child in wrapper)


def make_document(fragment, path):
    root = html.fragment_fromstring(fragment, create_parent='div')
    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'SimSun')
    pf = normal.paragraph_format
    pf.line_spacing = 1.4
    pf.first_line_indent = Pt(21)
    pf.space_after = Pt(8)

    def inline(parent, paragraph, bold=False, italic=False):
        def text(value):
            if value:
                r = paragraph.add_run(value)
                r.bold, r.italic = bold, italic
        text(parent.text)
        for child in parent:
            tag = etree.QName(child).localname.lower()
            if tag == 'math':
                source = etree.tostring(child, encoding='unicode', with_tail=False)
                try:
                    omml = mathml2omml.convert(source)
                except RuntimeError as exc:
                    if 'Too large number of children for ' not in str(exc):
                        raise
                    repaired = _group_excess_mathml_children(source)
                    if repaired == source:
                        raise
                    omml = mathml2omml.convert(repaired)
                omml = _hide_square_root_degree_placeholder(omml)
                # The converter returns an XML fragment without namespace declarations.
                wrapper = etree.fromstring(('<root xmlns:m="' + MATH_NS +
                    '" xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">' + omml + '</root>').encode())
                for equation in wrapper:
                    paragraph._p.append(equation)
            elif tag == 'br':
                paragraph.add_run().add_break()
            elif tag not in ('script', 'style', 'button'):
                inline(child, paragraph, bold or tag in ('b', 'strong'), italic or tag in ('i', 'em'))
            text(child.tail)

    blocks = {'p', 'div', 'ul', 'ol', 'li', 'blockquote', 'h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'pre', 'table'}
    heading_levels = {f'h{level}': level - 1 for level in range(1, 7)}

    def emit(node, container=doc):
        tag = etree.QName(node).localname.lower()
        if tag in ('script', 'style', 'button', 'hr'):
            return
        if tag == 'table':
            rows = node.xpath('./tr|./thead/tr|./tbody/tr|./tfoot/tr')
            if not rows:
                return
            cols = max(len(r.xpath('./td|./th')) for r in rows)
            table = container.add_table(rows=0, cols=cols)
            table.style = 'Table Grid'
            for row in rows:
                cells = table.add_row().cells
                for i, cell in enumerate(row.xpath('./td|./th')):
                    inline(cell, cells[i].paragraphs[0], bold=cell.tag == 'th')
                    cells[i].paragraphs[0].paragraph_format.first_line_indent = Pt(0)
            return
        if any(etree.QName(c).localname.lower() in blocks for c in node):
            if node.text and node.text.strip():
                container.add_paragraph(node.text)
            for child in node:
                emit(child, container)
                if child.tail and child.tail.strip():
                    container.add_paragraph(child.tail)
            return
        if not ''.join(node.itertext()).strip():
            return
        p = container.add_paragraph()
        # Lists intentionally become ordinary paragraphs: no Markdown bullet or dash.
        inline(node, p)
        if tag.startswith('h') and tag[1:].isdigit():
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_before = Pt(24 if tag == 'h1' else 10)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_with_next = True
            for run in p.runs:
                run.font.size = Pt(14 if tag == 'h1' else 13 if tag == 'h2' else 10.5)
                run.bold = True
        maths = node.xpath('.//math')
        if maths and all(m.get('display') == 'block' for m in maths):
            outside = ''.join(node.itertext())
            mathtext = ''.join(''.join(m.itertext()) for m in maths)
            if outside.strip() == mathtext.strip():
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.space_before = Pt(8)
        if tag == 'pre':
            p.paragraph_format.first_line_indent = Pt(0)
        if tag in heading_levels:
            outline_level = OxmlElement('w:outlineLvl')
            outline_level.set(qn('w:val'), str(heading_levels[tag]))
            p._p.get_or_add_pPr().append(outline_level)

    emit(root)
    # Direct formatting survives pasting into a document with a different Normal style.
    paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for paragraph in paragraphs:
        fmt = paragraph.paragraph_format
        fmt.line_spacing = 1.4
        if fmt.first_line_indent is None:
            fmt.first_line_indent = Pt(21)
        if fmt.space_after is None:
            fmt.space_after = Pt(8)
        for run in paragraph.runs:
            run.font.name = 'Times New Roman'
            run._r.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'SimSun')
            run.font.color.rgb = RGBColor(0, 0, 0)
            if run.font.size is None:
                run.font.size = Pt(10.5)
    old = doc.settings.element.find(qn('m:mathPr'))
    if old is not None:
        doc.settings.element.remove(old)
    doc.settings.element.append(parse_xml('<m:mathPr xmlns:m="' + MATH_NS + '"><m:mathFont m:val="Cambria Math"/></m:mathPr>'))
    doc.save(path)


_executor = ThreadPoolExecutor(max_workers=1, thread_name_prefix='office-clipboard')
_local = threading.local()


def _copy(fragment, target):
    import pythoncom
    import win32com.client
    if not getattr(_local, 'initialized', False):
        pythoncom.CoInitialize()
        _local.initialized = True
    progid = 'kwps.Application' if target == 'wps' else 'Word.Application'
    # Attach without changing the user's application settings or closing their app.
    try:
        app = win32com.client.GetActiveObject(progid)
    except pythoncom.com_error:
        try:
            app = win32com.client.DispatchEx(progid)
        except pythoncom.com_error as exc:
            raise RuntimeError('请先安装并打开 ' + ('WPS 文字' if target == 'wps' else 'Microsoft Word')) from exc
    # Keep the clipboard owner alive; never quit the user's Word/WPS instance.
    _local.app = app
    with tempfile.TemporaryDirectory(prefix='markdone-copy-') as folder:
        path = str(Path(folder) / 'selection.docx')
        make_document(fragment, path)
        document = app.Documents.Open(path, ReadOnly=True, AddToRecentFiles=False, Visible=False)
        try:
            document.Content.Copy()
        finally:
            document.Close(0)
    return True


def copy_native(fragment, target='word'):
    return _executor.submit(_copy, fragment, target).result(timeout=45)
